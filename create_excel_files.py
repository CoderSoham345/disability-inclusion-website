#!/usr/bin/env python3
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def style_header_row(ws, num_columns):
    """Apply styling to header row"""
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    
    for col in range(1, num_columns + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def style_data_rows(ws, num_rows, num_columns):
    """Apply styling to data rows"""
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for row in range(2, num_rows + 1):
        for col in range(1, num_columns + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            if row % 2 == 0:
                cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

def auto_adjust_columns(ws):
    """Auto-adjust column widths"""
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
        ws.column_dimensions[column_letter].width = adjusted_width

def create_excel_workbook(csv_file, sheet_name):
    """Create Excel workbook from CSV"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    with open(csv_file, 'r', encoding='utf-8') as f:
        reader = csv.reader(f)
        for row_idx, row in enumerate(reader, 1):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
    
    num_rows = ws.max_row
    num_columns = ws.max_column
    
    # Apply styling
    style_header_row(ws, num_columns)
    style_data_rows(ws, num_rows, num_columns)
    auto_adjust_columns(ws)
    
    # Freeze header row
    ws.freeze_panes = "A2"
    
    return wb

# Create individual workbooks
print("Creating Excel workbooks...")

# NGO Database
wb_ngo = create_excel_workbook('/vercel/share/v0-project/ngo_database.csv', 'NGO Database')
wb_ngo.save('/vercel/share/v0-project/Disability_Sports_NGO_Database.xlsx')
print("✓ NGO Database Excel created")

# Paralympians Database
wb_paralympians = create_excel_workbook('/vercel/share/v0-project/paralympians_database.csv', 'Paralympians')
wb_paralympians.save('/vercel/share/v0-project/Disability_Sports_Paralympians_Database.xlsx')
print("✓ Paralympians Database Excel created")

# Collaboration Database
wb_collab = create_excel_workbook('/vercel/share/v0-project/collaboration_opportunities.csv', 'Collaboration Partners')
wb_collab.save('/vercel/share/v0-project/Disability_Sports_Collaboration_Database.xlsx')
print("✓ Collaboration Database Excel created")

# Create combined workbook
print("\nCreating combined workbook...")
wb_combined = Workbook()
wb_combined.remove(wb_combined.active)  # Remove default sheet

# Add NGO sheet
with open('/vercel/share/v0-project/ngo_database.csv', 'r', encoding='utf-8') as f:
    reader = csv.reader(f)
    ws_ngo = wb_combined.create_sheet('NGO Database')
    for row_idx, row in enumerate(reader, 1):
        for col_idx, value in enumerate(row, 1):
            ws_ngo.cell(row=row_idx, column=col_idx, value=value)
    style_header_row(ws_ngo, ws_ngo.max_column)
    style_data_rows(ws_ngo, ws_ngo.max_row, ws_ngo.max_column)
    auto_adjust_columns(ws_ngo)
    ws_ngo.freeze_panes = "A2"

# Add Paralympians sheet
with open('/vercel/share/v0-project/paralympians_database.csv', 'r', encoding='utf-8') as f:
    reader = csv.reader(f)
    ws_paralympians = wb_combined.create_sheet('Paralympians')
    for row_idx, row in enumerate(reader, 1):
        for col_idx, value in enumerate(row, 1):
            ws_paralympians.cell(row=row_idx, column=col_idx, value=value)
    style_header_row(ws_paralympians, ws_paralympians.max_column)
    style_data_rows(ws_paralympians, ws_paralympians.max_row, ws_paralympians.max_column)
    auto_adjust_columns(ws_paralympians)
    ws_paralympians.freeze_panes = "A2"

# Add Collaboration sheet
with open('/vercel/share/v0-project/collaboration_opportunities.csv', 'r', encoding='utf-8') as f:
    reader = csv.reader(f)
    ws_collab = wb_combined.create_sheet('Collaboration Partners')
    for row_idx, row in enumerate(reader, 1):
        for col_idx, value in enumerate(row, 1):
            ws_collab.cell(row=row_idx, column=col_idx, value=value)
    style_header_row(ws_collab, ws_collab.max_column)
    style_data_rows(ws_collab, ws_collab.max_row, ws_collab.max_column)
    auto_adjust_columns(ws_collab)
    ws_collab.freeze_panes = "A2"

wb_combined.save('/vercel/share/v0-project/Disability_Sports_India_Complete_Database.xlsx')
print("✓ Combined workbook created")

print("\n✅ All Excel files created successfully!")
print("\nGenerated files:")
print("- Disability_Sports_NGO_Database.xlsx")
print("- Disability_Sports_Paralympians_Database.xlsx")
print("- Disability_Sports_Collaboration_Database.xlsx")
print("- Disability_Sports_India_Complete_Database.xlsx")
